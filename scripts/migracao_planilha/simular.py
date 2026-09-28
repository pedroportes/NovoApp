"""
Simulação da migração da planilha AppSheet -> FlowDrain (NÃO grava nada no banco).

Lê:
  - backup_migracao_20260928/*.json (retrato do banco)
  - backup_migracao_20260928/planilha_original/Novos.csv e Antigos.csv
Gera:
  - backup_migracao_20260928/simulacao.xlsx (para o Pedro conferir)
  - backup_migracao_20260928/plano.json (o que o passo de aplicar vai executar)
"""
import csv, json, os, re, sys, unicodedata
from collections import Counter, defaultdict

sys.stdout.reconfigure(encoding='utf-8')
BASE = r'C:\Users\pedro\NovoApp\backup_migracao_20260928'
EMPRESA = '58f0512e-8a00-4c31-ba32-f67f9b9ddcbe'

# Código da planilha (aba Empresas) -> id da marca no FlowDrain
MARCA_DA_PLANILHA = {
    '86676733': '93bd1248-9bcd-4e69-9d13-5569ae63db03',  # Curitibana
    '46a54a5d': '2c6b7aee-3453-49f6-88a5-f09879f8aafb',  # Hidro Curitiba
    '1639370a': '1372b5a0-b28e-4d8c-a1d6-75d6a47b0553',  # São José
    'be8edc77': 'fdf65712-50c5-4b71-a754-eac3a9eb4b5d',  # Nossa Cidade
    'Batel': 'e0de0000-0000-0000-0000-000000000004',     # Batel
    'Paulo': 'e0de0000-0000-0000-0000-000000000003',     # O Desentupidor
    '195ae451': '7fb4a63b-1dae-4f74-89bd-aa2e3bb917d4',  # Aqui Perto
}
TEC_JORGE_E_PEDRO = 'a0da0000-0000-0000-0000-000000000016'  # será criado no passo de aplicar
TEC_PEDRO_E_GRACA = 'a0da0000-0000-0000-0000-000000000001'
MARCA_O_DESENTUPIDOR = 'e0de0000-0000-0000-0000-000000000003'
FONE_TESTE = '41984501037'  # telefone do Pedro: clientes com ele são teste


def separa_dois_fones(d):
    """'99121174533481745' -> ['991211745', '33481745'] (dois números grudados)."""
    if not 16 <= len(d) <= 22:
        return None
    melhores = []
    for i in range(8, len(d) - 7):
        a, b = d[:i].lstrip('0'), d[i:].lstrip('0')
        ok = lambda x: (len(x) == 9 and x[0] == '9') or (len(x) == 8 and x[0] in '23456789') or (len(x) in (10, 11) and x[:2] in ('41', '42', '43', '44', '45', '46'))
        if ok(a) and ok(b):
            melhores.append(((len(a) == 9) + (len(b) == 9), a, b))
    if not melhores:
        return None
    melhores.sort(reverse=True)
    return [melhores[0][1], melhores[0][2]]


def load_json(nome):
    j = json.load(open(os.path.join(BASE, nome + '.json'), encoding='utf-8-sig'))
    rows = j['rows'][0] if isinstance(j, dict) else j
    return list(rows.values())[0] if isinstance(rows, dict) else rows


def load_csv(nome):
    return list(csv.DictReader(open(os.path.join(BASE, 'planilha_original', nome), encoding='utf-8-sig')))


def sem_acento(s):
    return unicodedata.normalize('NFD', s or '').encode('ascii', 'ignore').decode().lower().strip()


def so_dig(s):
    return re.sub(r'\D', '', str(s or ''))


# ---------------- Telefone: sempre (41) 99999-9999 / (41) 3333-4444 ----------------
def normaliza_fone(v):
    d = so_dig(v)
    if not d:
        return '', ''
    if len(d) in (12, 13) and d.startswith('55'):
        d = d[2:]
    if len(d) in (11, 12) and d.startswith('0'):
        d = d[1:]
    if len(d) in (8, 9):
        d = '41' + d
    if len(d) == 10 and d[2] in '6789':  # celular antigo sem o 9
        d = d[:2] + '9' + d[2:]
    if len(d) == 11:
        return d, f'({d[:2]}) {d[2:7]}-{d[7:]}'
    if len(d) == 10:
        return d, f'({d[:2]}) {d[2:6]}-{d[6:]}'
    return d, None  # inválido: não mexe


# ---------------- Endereço (mesma regra do app: clientSpreadsheet.splitAddress) ----------------
COMPLEMENTO_RE = r'^(casa|ap|apto|apartamento|bloco|bl|sala|sl|fundos|sobrado|lote|lt|loja|lj|andar|conj|cj|galp[aã]o)\b'
def split_endereco(full):
    s = re.sub(r'\s+', ' ', full or '').strip()
    s = re.sub(r',?\s*Brasil$', '', s, flags=re.I)
    out = {}
    m = re.search(r'\b(\d{2}\.?\d{3})-?(\d{3})\b', s)
    if m:
        out['cep'] = so_dig(m.group(0))
        s = s.replace(m.group(0), '')
        s = re.sub(r'\bCEP\b:?', '', s, flags=re.I)
    m = re.search(r'(?:\s-\s|,\s*|/)([A-Z]{2})\s*,?\s*$', s)
    if m:
        out['uf'] = m.group(1)
        s = s[:m.start()]
    parts = [p.strip() for p in re.split(r'\s*,\s*|\s+-\s+', s) if p.strip()]
    if not parts:
        return out
    rua = parts.pop(0)
    m = re.match(r'^(.*\D)\s+(\d{1,6}[A-Za-z]?)$', rua)
    if m and not re.match(r'^(br|pr|rod|km)$', (m.group(1).strip().split(' ') or [''])[-1], re.I):
        rua, out['numero'] = m.group(1).strip(), m.group(2)
    out['logradouro'] = rua
    if 'numero' not in out and parts and re.match(r'^(\d{1,6}[A-Za-z]?|s/?n)$', parts[0], re.I):
        out['numero'] = parts.pop(0)
    elif 'numero' not in out and parts:
        # "194 Rio Pequeno" -> nº 194 + bairro | "341 casa 3" -> nº 341 + complemento
        m = re.match(r'^(\d{1,6})\s+(.+)$', parts[0])
        if m:
            parts.pop(0)
            out['numero'] = m.group(1)
            resto = m.group(2).strip()
            if re.match(COMPLEMENTO_RE, resto, re.I):
                out['complemento'] = resto
            else:
                parts.insert(0, resto)
    if parts and re.match(COMPLEMENTO_RE, parts[0], re.I) and 'complemento' not in out:
        out['complemento'] = parts.pop(0)
    if len(parts) >= 3:
        out['complemento'] = parts.pop(0)
    if len(parts) >= 2:
        out['bairro'], out['cidade'] = parts[0], parts[1]
    elif len(parts) == 1:
        out['bairro'] = parts[0]
    return out


def endereco_quebrado(c):
    num = (c.get('numero') or '').strip()
    bairro = (c.get('bairro') or '').strip()
    rua = (c.get('logradouro') or '').strip()
    if re.fullmatch(r'0\d{2}', num) or re.fullmatch(r'\d+', bairro):
        return 'número/bairro trocados'
    if re.search(r'\bCEP\b', rua, re.I) or rua.count(',') >= 1:
        return 'rua com o endereço inteiro'
    if not num and re.search(r'\s\d{1,5}$', rua):
        return 'número junto da rua'
    if not bairro:
        return 'sem bairro'
    return ''


def primeiro_nome(n):
    t = re.sub(r'[^a-z0-9 ]', ' ', sem_acento(n)).split()
    return t[0] if t else ''


# ================= Carrega =================
clientes = load_json('clientes')
os_db = load_json('ordens_servico')
usuarios = {u['id']: u for u in load_json('usuarios')}
marcas = {m['id']: m for m in load_json('empresas_marcas')}
novos = [r for r in load_csv('Novos.csv') if r['Id'].strip() or r['Nome Cliente'].strip()]
# Linhas lançadas em dobro na planilha (mesmo Id, data, valor e cliente): o Pedro mandou manter uma só (28/09).
DOBRADAS = {'2e94bfd4', '5e6d6e1d'}  # Rosa 18/10/2024 e Vanessa 19/11/2024
_vistas = set()
_sem_dobra = []
for r in novos:
    k = (r['Id'].strip(), r['Data'], r['Valor Total'], r['Nome Cliente'].strip())
    if k[0] in DOBRADAS and k in _vistas:
        continue
    _vistas.add(k)
    _sem_dobra.append(r)
novos = _sem_dobra
antigos = [r for r in load_csv('Antigos.csv') if r['Cliente'].strip() or r['Celular'].strip()]
cli_by_id = {c['id']: c for c in clientes}
nome_tec = {v['nome']: k for k, v in usuarios.items() if v.get('nome')}

plano = {'clientes_update': {}, 'clientes_merge': [], 'os_update': {}, 'os_novas': [], 'antigos': [], 'comissoes_tecnico': {}}
rel = defaultdict(list)  # abas do relatório

# ================= 1. Telefones dos clientes =================
fone_invalido = 0
for c in clientes:
    orig = c.get('whatsapp') or c.get('telefone') or ''
    d, fmt = normaliza_fone(orig)
    upd = {}
    if orig and fmt is None:
        dois = separa_dois_fones(so_dig(orig))
        if dois:
            _, f1 = normaliza_fone(dois[0])
            _, f2 = normaliza_fone(dois[1])
            if f1 and f2:
                d, fmt = normaliza_fone(dois[0])
                ref = (c.get('referencia') or '').strip()
                upd['referencia'] = (ref + ' | ' if ref else '') + f'Outro telefone: {f2}'
                rel['Dois telefones separados'].append({'cliente': c['nome_razao'], 'antes': orig, 'WhatsApp': f1, 'outro telefone (vai na referência)': f2})
    if fmt and (c.get('whatsapp') != fmt or (c.get('telefone') or '') != fmt):
        upd['whatsapp'] = fmt
        upd['telefone'] = fmt
    elif orig and fmt is None:
        fone_invalido += 1
        rel['Telefones inválidos'].append({'cliente': c['nome_razao'], 'telefone': orig, 'id': c['id']})
    c['_fone'] = d if fmt else ''
    if upd:
        plano['clientes_update'].setdefault(c['id'], {}).update(upd)
        if len(rel['Telefones (amostra)']) < 300:
            rel['Telefones (amostra)'].append({'cliente': c['nome_razao'], 'antes': orig, 'depois': fmt})

# ================= 2. Endereços =================
for c in clientes:
    motivo = endereco_quebrado(c)
    full = (c.get('endereco') or c.get('address') or '').strip()
    if not motivo or not full:
        continue
    novo = split_endereco(full)
    upd = {}
    for campo in ('logradouro', 'numero', 'complemento', 'bairro', 'cidade', 'cep', 'uf'):
        v = (novo.get(campo) or '').strip()
        atual = (c.get(campo) or '').strip()
        if not v or v == atual:
            continue
        if motivo == 'sem bairro' and campo != 'bairro' and atual:
            continue  # só completa o bairro, não mexe no resto
        if campo == 'bairro':
            # só aceita bairro quando o endereço tinha bairro E cidade ("..., Bairro, Cidade")
            cidade_ref = sem_acento(novo.get('cidade') or c.get('cidade'))
            if not novo.get('cidade') or sem_acento(v) == cidade_ref or re.match(r'^\d', v):
                continue
        if campo == 'cidade' and atual:
            continue  # cidade já preenchida: mantém
        if campo == 'complemento' and atual:
            continue
        upd[campo] = v
    if upd:
        plano['clientes_update'].setdefault(c['id'], {}).update(upd)
        rel['Endereços corrigidos'].append({
            'cliente': c['nome_razao'], 'motivo': motivo, 'endereço original': full,
            'rua antes': c.get('logradouro'), 'nº antes': c.get('numero'), 'bairro antes': c.get('bairro'),
            'rua depois': upd.get('logradouro', c.get('logradouro')), 'nº depois': upd.get('numero', c.get('numero')),
            'bairro depois': upd.get('bairro', c.get('bairro')), 'cep depois': upd.get('cep', c.get('cep')),
        })

# ================= 3. Clientes duplicados (mesmo telefone) =================
os_por_cliente = Counter(o['cliente_id'] for o in os_db)
grupos = defaultdict(list)
for c in clientes:
    if c['_fone']:
        grupos[c['_fone']].append(c)
merge_de = {}  # id duplicado -> id que fica

# Clientes de teste (telefone do Pedro): serão apagados; as OS deles vão para conferência
teste_ids = {c['id'] for c in clientes if c['_fone'] == FONE_TESTE}
plano['clientes_apagar_teste'] = sorted(teste_ids)
plano['os_teste'] = {}
for o in os_db:
    if o.get('cliente_id') in teste_ids:
        real = 'translovato' in sem_acento(o.get('cliente_nome'))
        plano['os_teste'][o['id']] = 'manter' if real else 'apagar'
        rel['Teste - apagar'].append({
            'OS': o['id'][:8], 'data': (o.get('data_agendamento') or '')[:10], 'nome na OS': o.get('cliente_nome'),
            'cliente (teste)': cli_by_id[o['cliente_id']]['nome_razao'], 'valor': o.get('valor_total'), 'status': o.get('status'),
            'NFS-e': f"{o.get('nfe_numero') or ''} {o.get('nfe_status') or ''}".strip(),
            'sugestão': 'MANTER (cliente real com seu telefone): vai para um cliente próprio' if real else 'APAGAR',
        })
for cid in teste_ids:
    rel['Teste - apagar'].append({'OS': '—', 'nome na OS': '', 'cliente (teste)': cli_by_id[cid]['nome_razao'], 'sugestão': 'APAGAR cliente'})

for fone, g in grupos.items():
    if len(g) < 2 or fone == FONE_TESTE:
        continue
    g.sort(key=lambda c: c.get('created_at') or '')
    fica = g[0]
    for outro in g[1:]:
        mesmo_nome = primeiro_nome(outro['nome_razao']) == primeiro_nome(fica['nome_razao']) or \
            sem_acento(fica['nome_razao'])[:6] in sem_acento(outro['nome_razao']) or \
            sem_acento(outro['nome_razao'])[:6] in sem_acento(fica['nome_razao'])
        rua_a = re.sub(r'[^a-z0-9]', '', sem_acento(fica.get('logradouro')))[-10:]
        rua_b = re.sub(r'[^a-z0-9]', '', sem_acento(outro.get('logradouro')))[-10:]
        mesmo_end = bool(rua_a) and rua_a == rua_b and so_dig(fica.get('numero')) == so_dig(outro.get('numero'))
        linha = {
            'telefone': normaliza_fone(fone)[1], 'fica': fica['nome_razao'], 'fica endereço': f"{fica.get('logradouro') or ''}, {fica.get('numero') or ''}",
            'fica OS': os_por_cliente[fica['id']], 'sai': outro['nome_razao'],
            'sai endereço': f"{outro.get('logradouro') or ''}, {outro.get('numero') or ''}", 'sai OS': os_por_cliente[outro['id']],
            'sai criado em': (outro.get('created_at') or '')[:10],
        }
        if mesmo_nome or mesmo_end:
            merge_de[outro['id']] = fica['id']
            completar = {}
            for campo in ('cpf_cnpj', 'email', 'cep', 'logradouro', 'numero', 'complemento', 'bairro', 'cidade', 'uf', 'referencia', 'marca_id'):
                if not (fica.get(campo) or '').strip() and (outro.get(campo) or '').strip():
                    completar[campo] = outro[campo]
            # nome: se o do duplicado for mais completo e começar igual, usa ele
            if len(outro['nome_razao'] or '') > len(fica['nome_razao'] or '') and sem_acento(outro['nome_razao']).startswith(sem_acento(fica['nome_razao'])):
                completar['nome_razao'] = outro['nome_razao']
            plano['clientes_merge'].append({'fica': fica['id'], 'sai': outro['id'], 'completar': completar})
            linha['motivo'] = 'mesmo nome' if mesmo_nome else 'mesmo endereço'
            rel['Clientes que serão juntados'].append(linha)
        else:
            rel['Mesmo telefone, NÃO junta'].append(linha)

# Mesmo cliente cadastrado de novo a cada chamado (sem telefone): mesmo nome + mesma rua + mesmo número
def chave_nome_end(c):
    nome = re.sub(r'[^a-z0-9]', '', sem_acento(re.sub(r'\(.*?\)', '', c['nome_razao'] or '')))
    rua = re.sub(r'^(rua|r|avenida|av|travessa|tv|alameda|al)\.?\s+', '', sem_acento(c.get('logradouro')))
    rua = re.sub(r'[^a-z0-9]', '', rua)
    num = so_dig(c.get('numero'))
    return f'{nome}|{rua}|{num}' if nome and rua and num and nome not in ('clientesemnome', 'semnome') else None


grupos_ne = defaultdict(list)
for c in clientes:
    if c['id'] in merge_de or c['id'] in teste_ids:
        continue
    k = chave_nome_end(c)
    if k:
        grupos_ne[k].append(c)
for k, g in grupos_ne.items():
    if len(g) < 2:
        continue
    g.sort(key=lambda c: c.get('created_at') or '')
    fica = g[0]
    while fica['id'] in merge_de:
        fica = cli_by_id[merge_de[fica['id']]]
    for outro in g[1:]:
        if outro['id'] == fica['id'] or outro['id'] in merge_de:
            continue
        fones = {fica['_fone'], outro['_fone']} - {''}
        if len(fones) > 1:
            continue  # telefones diferentes: pode ser outra pessoa na mesma casa, não junta
        merge_de[outro['id']] = fica['id']
        completar = {campo: outro[campo] for campo in ('whatsapp', 'telefone', 'cpf_cnpj', 'email', 'cep', 'complemento', 'bairro', 'referencia', 'marca_id')
                     if not (fica.get(campo) or '').strip() and (outro.get(campo) or '').strip()}
        if 'whatsapp' in completar and outro['_fone']:
            completar['whatsapp'] = completar['telefone'] = normaliza_fone(outro['_fone'])[1]
        plano['clientes_merge'].append({'fica': fica['id'], 'sai': outro['id'], 'completar': completar})
        rel['Clientes que serão juntados'].append({
            'telefone': normaliza_fone(fica['_fone'] or outro['_fone'])[1] if (fica['_fone'] or outro['_fone']) else '',
            'fica': fica['nome_razao'], 'fica endereço': f"{fica.get('logradouro') or ''}, {fica.get('numero') or ''}", 'fica OS': os_por_cliente[fica['id']],
            'sai': outro['nome_razao'], 'sai endereço': f"{outro.get('logradouro') or ''}, {outro.get('numero') or ''}", 'sai OS': os_por_cliente[outro['id']],
            'sai criado em': (outro.get('created_at') or '')[:10], 'motivo': 'mesmo nome e endereço (novo cadastro a cada chamado)',
        })

def destino(cid):
    """Segue a cadeia de junções até o cadastro que fica."""
    vistos = set()
    while cid in merge_de and cid not in vistos:
        vistos.add(cid)
        cid = merge_de[cid]
    return cid


for m in plano['clientes_merge']:
    m['fica'] = destino(m['fica'])

# ================= 4. OS: casa planilha x banco =================
def money(s):
    s = re.sub(r'[^\d,.-]', '', s or '')
    if not s:
        return 0.0
    if ',' in s:
        s = s.replace('.', '').replace(',', '.')
    try:
        return round(float(s), 2)
    except ValueError:
        return 0.0


def data_iso(s):
    m = re.match(r'(\d{2})/(\d{2})/(\d{4})', s or '')
    return f'{m[3]}-{m[2]}-{m[1]}' if m else ''


idx = defaultdict(list)
for o in os_db:
    idx[((o.get('data_agendamento') or '')[:10], round(float(o.get('valor_total') or 0), 2))].append(o)
usados = set()
sem_par = []
for r in novos:
    chave = (data_iso(r['Data']), money(r['Valor Total']))
    cands = [o for o in idx.get(chave, []) if o['id'] not in usados]
    if not cands:
        sem_par.append(r)
        continue
    pn = primeiro_nome(r['Nome Cliente'])
    marca_pl = MARCA_DA_PLANILHA.get(r['Empresa Prestadora'].strip())
    tec_pl = r['Tecnicofez'].strip()

    def pontos(c):
        return (2 if pn and primeiro_nome(c.get('cliente_nome')) == pn else 0) + \
               (1 if marca_pl and c.get('marca_id') == marca_pl else 0) + \
               (1 if tec_pl and usuarios.get(c.get('tecnico_id'), {}).get('nome') == tec_pl else 0)
    o = max(cands, key=pontos)
    confianca = 'nome+data+valor' if pn and primeiro_nome(o.get('cliente_nome')) == pn else 'só data+valor'
    usados.add(o['id'])
    upd = {'origem_id': r['Id'].strip() or None}

    tec = r['Tecnicofez'].strip()
    tec_atual = usuarios.get(o.get('tecnico_id'), {}).get('nome')
    novo_tec = {'Jorge e Pedro': TEC_JORGE_E_PEDRO, 'Pedro e Graça': TEC_PEDRO_E_GRACA}.get(tec)
    if novo_tec and o.get('tecnico_id') != novo_tec:
        upd['tecnico_id'] = novo_tec
        plano['comissoes_tecnico'][o['id']] = novo_tec
        rel['OS: técnico corrigido'].append({'OS': o['id'][:8], 'data': chave[0], 'cliente': o.get('cliente_nome'), 'antes': tec_atual, 'depois': tec})
    elif tec and tec != tec_atual and not novo_tec:
        rel['OS: técnico diferente (não mexe)'].append({'OS': o['id'][:8], 'data': chave[0], 'cliente': o.get('cliente_nome'), 'na planilha': tec, 'no app': tec_atual})
    tec_final = tec if novo_tec else tec_atual

    marca_certa = MARCA_DA_PLANILHA.get(r['Empresa Prestadora'].strip())
    # O Desentupidor (empresa do Paulo): só vai para lá se quem fez foi o Paulo; se foi Pedro / Pedro e Graça fica onde está
    if marca_certa == MARCA_O_DESENTUPIDOR and tec_final != 'Paulo':
        if o.get('marca_id') != marca_certa:
            rel['OS: empresa mantida'].append({'OS': o['id'][:8], 'data': chave[0], 'cliente': o.get('cliente_nome'), 'valor': chave[1], 'técnico': tec_final,
                                              'empresa na planilha': 'O Desentupidor', 'fica em': marcas.get(o.get('marca_id'), {}).get('nome') or '(sem empresa)'})
        marca_certa = None
    if marca_certa and o.get('marca_id') != marca_certa:
        upd['marca_id'] = marca_certa
        rel['OS: empresa corrigida'].append({'OS': o['id'][:8], 'data': chave[0], 'cliente': o.get('cliente_nome'), 'valor': chave[1], 'técnico': tec_final,
                                            'antes': marcas.get(o.get('marca_id'), {}).get('nome'), 'depois': marcas[marca_certa]['nome']})

    d, fmt = normaliza_fone(o.get('cliente_whatsapp'))
    if fmt and o.get('cliente_whatsapp') != fmt:
        upd['cliente_whatsapp'] = fmt
    if merge_de.get(o.get('cliente_id')):
        upd['cliente_id'] = destino(o['cliente_id'])
    if o.get('cliente_id') in teste_ids and plano['os_teste'].get(o['id']) == 'manter':
        # cliente real que estava com o telefone do Pedro: ganha cadastro próprio (sem o telefone de teste)
        end = split_endereco(r['Endereco'])
        nome = r['Nome Cliente'].strip() + (f" ({r['Emp.Cond'].strip()})" if r['Emp.Cond'].strip() else '')
        plano.setdefault('clientes_criar', {})[o['id']] = {
            'nome_razao': nome, 'endereco': r['Endereco'].strip(), 'marca_id': upd.get('marca_id') or o.get('marca_id'),
            **{k: v for k, v in end.items() if k in ('logradouro', 'numero', 'complemento', 'bairro', 'cidade', 'uf', 'cep')},
        }
    # mídia da planilha (será copiada do Drive numa etapa própria)
    midias = {k: r[k].strip() for k in ('Assinatura', 'Imagem', 'Imagem 2', 'Video') if r[k].strip()}
    if midias:
        upd['_midias'] = midias
    if confianca != 'nome+data+valor':
        rel['OS: casamento a conferir'].append({'OS': o['id'][:8], 'data': chave[0], 'valor': chave[1], 'planilha': r['Nome Cliente'], 'app': o.get('cliente_nome')})
    plano['os_update'][o['id']] = upd

# OS do banco que viraram de cliente juntado mas não estão na planilha
for o in os_db:
    if o['id'] not in plano['os_update'] and merge_de.get(o.get('cliente_id')):
        plano['os_update'][o['id']] = {'cliente_id': destino(o['cliente_id'])}

for r in sem_par:
    item = {'Id': r['Id'], 'data': r['Data'], 'cliente': r['Nome Cliente'], 'valor': r['Valor Total'],
            'empresa': r['Empresa Prestadora'], 'técnico': r['Tecnicofez'], 'serviço': r['Serviço Feito']}
    if money(r['Valor Total']) > 0 and r['Nome Cliente'].strip():
        plano['os_novas'].append(r)
        rel['OS novas (entram)'].append(item)
    else:
        rel['OS da planilha que NÃO entram'].append({**item, 'motivo': 'sem valor ou sem nome'})

# ================= 5. Antigos (2015–2021) -> histórico leve =================
idx_fone = {}
idx_nome_rua = {}
for c in clientes:
    if c['id'] in merge_de or c['id'] in teste_ids:
        continue
    if c['_fone']:
        idx_fone.setdefault(c['_fone'], c['id'])
    k = primeiro_nome(c['nome_razao']) + '|' + re.sub(r'[^a-z0-9]', '', sem_acento(c.get('logradouro')))[:12]
    idx_nome_rua.setdefault(k, c['id'])
ligados = 0
for r in antigos:
    d, fmt = normaliza_fone(r['Celular'])
    cid = idx_fone.get(d) if fmt else None
    if not cid:
        end = split_endereco(r['Endereço'])
        k = primeiro_nome(r['Cliente']) + '|' + re.sub(r'[^a-z0-9]', '', sem_acento(end.get('logradouro') or r['Endereço']))[:12]
        cid = idx_nome_rua.get(k)
    ligados += bool(cid)
    plano['antigos'].append({
        'origem_id': r['Id'], 'cliente_id': cid, 'cliente_nome': r['Cliente'].strip(), 'telefone': fmt or r['Celular'].strip(),
        'endereco': r['Endereço'].strip(), 'cidade': r['Cidade'].strip(), 'data': data_iso(r['Data']) or None,
        'servico': r['Serviço realizado'].strip(), 'valor': money(r['Valor']), 'pagamento': r['Meio de pagamento'].strip(),
        'quem_fez': r['Quem fez'].strip(), 'detalhes': r['Maiores detalhes'].strip(),
    })
    if not cid and len(rel['Antigos sem cliente (amostra)']) < 300:
        rel['Antigos sem cliente (amostra)'].append({'Id': r['Id'], 'cliente': r['Cliente'], 'telefone': r['Celular'], 'endereço': r['Endereço'], 'data': r['Data']})

# ================= Resumo =================
n_midia = sum(1 for u in plano['os_update'].values() if '_midias' in u)
resumo = [
    ('Clientes no app hoje', len(clientes)),
    ('Clientes com telefone ajustado para o formato (41) 99999-9999', sum(1 for u in plano['clientes_update'].values() if 'whatsapp' in u)),
    ('Telefones inválidos (não mexe, aparecem na aba)', fone_invalido),
    ('Clientes com endereço corrigido', len(rel['Endereços corrigidos'])),
    ('Clientes duplicados que serão juntados', len(plano['clientes_merge'])),
    ('Mesmo telefone mas pessoas diferentes (não junta)', len(rel['Mesmo telefone, NÃO junta'])),
    ('Clientes com dois telefones separados (2º vai na referência)', len(rel['Dois telefones separados'])),
    ('Clientes de teste (seu telefone) que serão apagados', len(teste_ids)),
    ('OS desses clientes de teste que serão apagadas', sum(1 for v in plano['os_teste'].values() if v == 'apagar')),
    ('OS desses clientes que são reais e ficam (ganham cliente próprio)', sum(1 for v in plano['os_teste'].values() if v == 'manter')),
    ('Clientes depois da junção', len(clientes) - len(plano['clientes_merge']) - len(teste_ids) + len(plano.get('clientes_criar', {}))),
    ('OS da planilha (Novos)', len(novos)),
    ('OS achadas no app', len(usados)),
    ('OS a conferir (achadas só por data+valor)', len(rel['OS: casamento a conferir'])),
    ('OS com empresa corrigida', len(rel['OS: empresa corrigida'])),
    ('OS com técnico corrigido (Jorge e Pedro / Pedro e Graça)', len(rel['OS: técnico corrigido'])),
    ('OS da O Desentupidor feitas por Pedro/Pedro e Graça (ficam onde estão)', len(rel['OS: empresa mantida'])),
    ('OS com técnico diferente mas que ficam como estão', len(rel['OS: técnico diferente (não mexe)'])),
    ('OS novas que vão entrar', len(plano['os_novas'])),
    ('OS da planilha que não entram (sem valor/nome)', len(rel['OS da planilha que NÃO entram'])),
    ('OS com assinatura/foto para copiar do Drive', n_midia),
    ('Serviços antigos (2015–2021) no histórico', len(plano['antigos'])),
    ('   ligados a um cliente', ligados),
    ('   sem cliente achado', len(plano['antigos']) - ligados),
]
for k, v in resumo:
    print(f'{k}: {v}')

json.dump(plano, open(os.path.join(BASE, 'plano.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)

# ================= Planilha de conferência =================
sys.path.insert(0, os.path.dirname(__file__))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook()
ws = wb.active
ws.title = 'Resumo'
ws.append(['O que vai acontecer', 'Quantidade'])
for k, v in resumo:
    ws.append([k, v])
ws.column_dimensions['A'].width = 70
ws.column_dimensions['B'].width = 14
ordem = ['Teste - apagar', 'Clientes que serão juntados', 'Mesmo telefone, NÃO junta', 'Endereços corrigidos', 'Dois telefones separados',
         'Telefones (amostra)', 'Telefones inválidos', 'OS: empresa corrigida', 'OS: empresa mantida', 'OS: técnico corrigido', 'OS: técnico diferente (não mexe)', 'OS: casamento a conferir',
         'OS novas (entram)', 'OS da planilha que NÃO entram', 'Antigos sem cliente (amostra)']
for nome in ordem:
    linhas = rel.get(nome) or []
    w = wb.create_sheet(nome[:31].replace(':', ' -'))
    if not linhas:
        w.append(['(nada)'])
        continue
    cab = list(linhas[0].keys())
    w.append(cab)
    for l in linhas:
        w.append([str(l.get(k) if l.get(k) is not None else '') for k in cab])
    for i, _ in enumerate(cab):
        w.column_dimensions[chr(65 + i) if i < 26 else 'Z'].width = 24
for w in wb.worksheets:
    for cell in w[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='0F766E')
        cell.alignment = Alignment(wrap_text=True)
    w.freeze_panes = 'A2'
out = os.path.join(BASE, 'simulacao.xlsx')
wb.save(out)
print('\nPlanilha:', out)
